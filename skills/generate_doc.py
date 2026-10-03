import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_term_document(output_path):
    doc = docx.Document()

    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Helper styling functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def add_heading_styled(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Dark Blue
        elif level == 2:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue
        return p

    # --- DOCUMENT HEADER TITLE ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("BẢNG GIẢI THÍCH THUẬT NGỮ CHUYÊN MÔN")
    r_title.font.name = 'Calibri'
    r_title.font.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # --- METADATA TABLE (Term, Acronym, Topic) ---
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    meta_data = [
        ("Term (Thuật ngữ)", "Retrieval-Augmented Generation"),
        ("Acronym (Viết tắt)", "RAG"),
        ("Topic (Chủ đề)", "Artificial Intelligence & Natural Language Processing")
    ]

    col_widths = [Inches(2.0), Inches(4.5)]

    for idx, (label, val) in enumerate(meta_data):
        row = table.rows[idx]
        
        # Cell 0 (Label)
        cell_lbl = row.cells[0]
        cell_lbl.width = col_widths[0]
        set_cell_background(cell_lbl, "1E3A8A") # Navy Blue
        set_cell_margins(cell_lbl, top=120, bottom=120, left=150, right=150)
        p0 = cell_lbl.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(label)
        r0.font.name = 'Calibri'
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        # Cell 1 (Value)
        cell_val = row.cells[1]
        cell_val.width = col_widths[1]
        set_cell_background(cell_val, "F8FAFC") # Light Slate
        set_cell_margins(cell_val, top=120, bottom=120, left=150, right=150)
        p1 = cell_val.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(val)
        r1.font.name = 'Calibri'
        r1.font.bold = (idx == 0 or idx == 1)
        r1.font.size = Pt(11.5 if idx == 0 else 11)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION 1: DIỄN GIẢI NGẮN ---
    add_heading_styled("1. Diễn giải ngắn (Short Summary)", level=1)

    # Callout box for short explanation
    tbl_summary = doc.add_table(rows=1, cols=1)
    tbl_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_sum = tbl_summary.rows[0].cells[0]
    c_sum.width = Inches(6.5)
    set_cell_background(c_sum, "EFF6FF") # Light Blue tint
    set_cell_margins(c_sum, top=140, bottom=140, left=180, right=180)

    # Left border styling for callout box
    tcPr = c_sum._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p_sum = c_sum.paragraphs[0]
    p_sum.paragraph_format.space_after = Pt(0)
    p_sum.paragraph_format.line_spacing = 1.15
    r_sum = p_sum.add_run(
        "Retrieval-Augmented Generation (RAG) là kỹ thuật kết hợp giữa khả năng truy xuất "
        "thông tin từ cơ sở dữ liệu bên ngoài và khả năng sáng tạo ngôn ngữ của các mô hình AI lớn (LLM). "
        "Thay vì chỉ dựa vào tri thức sẵn có đã được huấn luyện, AI sẽ tìm kiếm thông tin chính xác từ tài liệu "
        "riêng của bạn trước khi tổng hợp thành câu trả lời."
    )
    r_sum.font.name = 'Calibri'
    r_sum.font.size = Pt(11)
    r_sum.font.italic = True
    r_sum.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # --- SECTION 2: ĐỊNH NGHĨA CHI TIẾT ---
    add_heading_styled("2. Định nghĩa chi tiết (Definition)", level=1)

    p_def1 = doc.add_paragraph()
    p_def1.paragraph_format.space_after = Pt(6)
    p_def1.paragraph_format.line_spacing = 1.15
    r_def1 = p_def1.add_run(
        "RAG (Retrieval-Augmented Generation) là một kiến trúc tiên tiến trong Xử lý Ngôn ngữ Tự nhiên (NLP), "
        "giải quyết hạn chế cốt lõi của các mô hình AI lớn bằng cách chia quy trình xử lý thành hai giai đoạn chính:"
    )
    r_def1.font.name = 'Calibri'
    r_def1.font.size = Pt(11)
    r_def1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    items = [
        (
            "Giai đoạn Truy xuất (Retrieval): ",
            "Khi nhận câu hỏi từ người dùng, hệ thống sử dụng thuật toán tìm kiếm ngữ nghĩa (semantic search) "
            "hoặc cơ sở dữ liệu vector (vector database) để truy xuất các đoạn văn bản (chunks) chứa thông tin liên quan nhất "
            "từ kho dữ liệu tri thức bên ngoài (ví dụ: tài liệu nội bộ, file PDF, cơ sở dữ liệu doanh nghiệp)."
        ),
        (
            "Giai đoạn Tạo câu trả lời (Generation): ",
            "Các đoạn dữ liệu truy xuất được kết hợp làm bối cảnh (context) cùng câu hỏi ban đầu đưa vào mô hình ngôn ngữ lớn "
            "(như GPT-4, Gemini). Mô hình sẽ phân tích bối cảnh này để tổng hợp ra câu trả lời chuẩn xác và tự nhiên."
        )
    ]

    for title, desc in items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(4)
        p_item.paragraph_format.line_spacing = 1.15
        
        r_t = p_item.add_run(title)
        r_t.font.name = 'Calibri'
        r_t.font.bold = True
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        
        r_d = p_item.add_run(desc)
        r_d.font.name = 'Calibri'
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    add_heading_styled("Lợi ích cốt lõi của RAG:", level=2)

    benefits = [
        (
            "Giảm hiện tượng ảo giác (Hallucination): ",
            "Buộc mô hình AI phải đưa ra câu trả lời dựa trên dữ liệu thực tế được cung cấp, tránh tình trạng tự suy đoán thông tin sai lệch."
        ),
        (
            "Cập nhật tri thức theo thời gian thực: ",
            "Dữ liệu mới có thể được bổ sung ngay lập tức vào kho lưu trữ mà không cần tốn chi phí và thời gian huấn luyện lại (fine-tune) mô hình."
        ),
        (
            "Bảo mật thông tin doanh nghiệp: ",
            "Cho phép tận dụng sức mạnh AI trên tài liệu nội bộ mà không cần tải dữ liệu nhạy cảm lên các mô hình công cộng để huấn luyện."
        )
    ]

    for b_title, b_desc in benefits:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.left_indent = Inches(0.25)
        p_b.paragraph_format.space_after = Pt(4)
        p_b.paragraph_format.line_spacing = 1.15
        
        r_bt = p_b.add_run(b_title)
        r_bt.font.name = 'Calibri'
        r_bt.font.bold = True
        r_bt.font.size = Pt(11)
        r_bt.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        r_bd = p_b.add_run(b_desc)
        r_bd.font.name = 'Calibri'
        r_bd.font.size = Pt(11)
        r_bd.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # --- SECTION 3: VÍ DỤ THỰC TẾ ---
    add_heading_styled("3. Ví dụ thực tế (Example)", level=1)

    p_ex_intro = doc.add_paragraph()
    p_ex_intro.paragraph_format.space_after = Pt(6)
    p_ex_intro.paragraph_format.line_spacing = 1.15
    r_ex_i = p_ex_intro.add_run(
        "Tình huống: Một doanh nghiệp bảo hiểm triển khai Chatbot AI để hỗ trợ nhân viên tư vấn khách hàng. "
        "Nhân viên hỏi: \"Hợp đồng gói Vàng có bồi thường cho chi phí điều trị tai nạn thể thao không?\""
    )
    r_ex_i.font.name = 'Calibri'
    r_ex_i.font.size = Pt(11)
    r_ex_i.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    tbl_ex = doc.add_table(rows=2, cols=2)
    tbl_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ex.autofit = False

    ex_headers = ["Kịch bản KHÔNG dùng RAG", "Kịch bản CÓ ứng dụng RAG"]
    ex_contents = [
        "AI chỉ trả lời dựa trên tri thức chung đã học từ trước. Nó có thể đưa ra câu trả lời chung chung hoặc suy đoán sai chính xác điều khoản của công ty (Ví dụ: 'Thông thường các gói bảo hiểm sức khỏe đều hỗ trợ...').",
        "1. Retrieval: Hệ thống tự động truy xuất kho 500 trang quy tắc bảo hiểm của công ty và tìm đúng Điều 12.3 của gói Vàng.\n\n2. Generation: AI tổng hợp từ Điều 12.3 và trả lời: 'Theo Điều 12.3 của Quy tắc Bảo hiểm gói Vàng, chi phí điều trị tai nạn thể thao không chuyên nghiệp được bồi thường tối đa 50 triệu VNĐ/vụ.'"
    ]

    for c_idx in range(2):
        # Header Cell
        cell_h = tbl_ex.rows[0].cells[c_idx]
        cell_h.width = Inches(3.25)
        set_cell_background(cell_h, "DC2626" if c_idx == 0 else "16A34A") # Red for no RAG, Green for RAG
        set_cell_margins(cell_h, top=100, bottom=100, left=120, right=120)
        p_h = cell_h.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_h = p_h.add_run(ex_headers[c_idx])
        r_h.font.name = 'Calibri'
        r_h.font.bold = True
        r_h.font.size = Pt(11)
        r_h.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Content Cell
        cell_c = tbl_ex.rows[1].cells[c_idx]
        cell_c.width = Inches(3.25)
        set_cell_background(cell_c, "FEF2F2" if c_idx == 0 else "F0FDF4")
        set_cell_margins(cell_c, top=120, bottom=120, left=140, right=140)
        p_c = cell_c.paragraphs[0]
        p_c.paragraph_format.line_spacing = 1.15
        r_c = p_c.add_run(ex_contents[c_idx])
        r_c.font.name = 'Calibri'
        r_c.font.size = Pt(10.5)
        r_c.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.save(output_path)
    print(f"File created successfully at: {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.getcwd(), "skills", "giai_thich_thuat_ngu_RAG.docx")
    create_term_document(out_file)
